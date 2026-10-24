"""Builds the project report (.docx) using the real numbers from outputs/metrics.json."""
import json

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

M = json.load(open("outputs/metrics.json"))
FIG = "outputs/figures/"
ACCENT = RGBColor(0x4B, 0x2E, 0xC9)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(2.3))
sec.top_margin = sec.bottom_margin = Cm(2.2)

st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
st.paragraph_format.space_after = Pt(6)
st.paragraph_format.line_spacing = 1.15
for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
    h = doc.styles[name]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.bold = True
    h.font.color.rgb = ACCENT
    h.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    h.paragraph_format.space_after = Pt(6)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para(text, bold=False, italic=False, align=None, size=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    return p


def bullets(items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def numbered(items):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph(f"{i}. {it}")
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.space_after = Pt(2)


def table(header, rows, widths=None, font=10):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold = True
        r.font.size = Pt(font)
        r.font.color.rgb = RGBColor(255, 255, 255)
        shade(c, "4B2EC9")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(font)
            if ri % 2 == 1:
                shade(cells[i], "F1EEFC")
    for row in t.rows[:-1]:
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


fig_no = [0]


def figure(file, caption, width=14):
    fig_no[0] += 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(FIG + file, width=Cm(width))
    c = para(f"Figure {fig_no[0]}: {caption}", italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=9.5)
    c.paragraph_format.space_after = Pt(10)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------------------------------ cover
for _ in range(5):
    doc.add_paragraph()
para("SkillOrbit", bold=True, size=16, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Machine Learning Capstone Project", size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_paragraph()
para("Student Performance Prediction System", bold=True, size=30, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Project Report", size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
for _ in range(4):
    doc.add_paragraph()
table(["Field", "Details"], [
    ["Submitted by", "[Your Name]"],
    ["Roll / Registration No.", "[Your Roll No.]"],
    ["Programme", "Machine Learning Capstone"],
    ["Institution", "SkillOrbit, HSR Layout, Bangalore"],
    ["Date of submission", "October 2026"],
    ["GitHub repository", "[Paste repository link]"],
    ["Deployment link", "[Paste Streamlit Cloud link]"],
], widths=[5, 10], font=11)
page_break()

# ------------------------------------------------------------------------------------------ contents
doc.add_heading("Table of Contents", 1)
toc = ["1. Introduction", "2. Objectives", "3. Industry Relevance", "4. System Design and Tech Stack",
       "5. Dataset Description", "6. Data Preprocessing (Module 1)", "7. Exploratory Data Analysis (Module 2)",
       "8. Model Development (Module 3)", "9. Results and Evaluation", "10. Performance Dashboard (Module 4)",
       "11. Recommendation System (Module 5)", "12. Testing", "13. Deployment",
       "14. Limitations and Future Work", "15. Conclusion", "16. References", "Appendix: Project Structure"]
for t in toc:
    doc.add_paragraph(t)
page_break()

# ------------------------------------------------------------------------------------------ 1 intro
doc.add_heading("1. Introduction", 1)
para("Academic performance is influenced by many factors: how regularly a student attends class, how many hours "
     "they study, how well they have done in previous assessments, and how actively they take part in learning. "
     "Institutions usually notice that a student is struggling only after the final results are out, which is too "
     "late to help.")
para("The Student Performance Prediction System is a machine learning application that analyzes student academic "
     "data, predicts the expected performance level (Low, Average or High) together with a numeric final score, "
     "identifies the factors that matter most, shows the results on an interactive dashboard, and generates "
     "personalised recommendations for improvement.")
para("The project follows an end-to-end machine learning workflow: dataset collection, data cleaning, exploratory "
     "analysis, model development, prediction generation, dashboard visualisation and recommendation generation.")

doc.add_heading("2. Objectives", 1)
bullets(["Analyze student academic data.",
         "Predict student performance levels (Low, Average, High) and the final score.",
         "Identify the key factors affecting performance.",
         "Generate performance insights and improvement recommendations.",
         "Visualize student analytics through an interactive dashboard and deploy it online."])

doc.add_heading("3. Industry Relevance", 1)
para("Prediction of student outcomes is a standard use case in EdTech platforms, learning management systems, "
     "academic analytics platforms and online learning applications. Early-warning systems built on such "
     "models allow teachers and mentors to identify at-risk students early, plan remedial classes and allocate "
     "support where it is most needed. This project reproduces a simplified version of that workflow and "
     "demonstrates how machine learning supports educational decision-making.")

# ------------------------------------------------------------------------------------------ 4 design
doc.add_heading("4. System Design and Tech Stack", 1)
doc.add_heading("4.1 Workflow", 2)
numbered(["Dataset collection", "Data cleaning", "Exploratory data analysis", "Model development",
          "Prediction generation", "Dashboard visualisation", "Recommendation generation"])
doc.add_heading("4.2 Technology stack", 2)
table(["Layer", "Tools used"], [
    ["Programming language", "Python 3"],
    ["Libraries", "Pandas, NumPy, Scikit-learn, Matplotlib, Seaborn, Joblib"],
    ["Dashboard", "Streamlit"],
    ["Model persistence", "Joblib (model_bundle.joblib)"],
    ["Deployment", "Streamlit Community Cloud"],
    ["Version control", "Git and GitHub"],
], widths=[5, 10])
doc.add_heading("4.3 Code organisation", 2)
para("The code is split so that the same logic is reused by training and by the dashboard: ml_core.py holds the "
     "cleaning, encoding, performance-level, prediction and recommendation functions; train_model.py runs the "
     "complete pipeline and saves the model and charts; app.py is the Streamlit dashboard.")

# ------------------------------------------------------------------------------------------ 5 dataset
doc.add_heading("5. Dataset Description", 1)
para("The project allows Kaggle datasets or self-created academic datasets. A self-created dataset was used, "
     "generated by generate_dataset.py with a fixed random seed so the results are reproducible. The dataset is "
     "synthetic: relationships between factors and the final score were built in, and noise, missing values and "
     "duplicate rows were added deliberately so that the preprocessing stage is realistic. Because the data is "
     "synthetic, the accuracy figures in this report describe how well the pipeline works on this dataset and "
     "should not be read as accuracy on real students.", italic=False)
para(f"The raw file contains {M['raw_rows']:,} rows and {M['raw_cols']} columns.")
table(["Column", "Type", "Description"], [
    ["student_id", "ID", "Unique student identifier"],
    ["gender, parental_education", "Categorical", "Student background"],
    ["internet_access, extracurricular, tutoring", "Categorical (Yes/No)", "Learning environment"],
    ["attendance", "Numeric (%)", "Class attendance percentage"],
    ["study_hours", "Numeric", "Self-study hours per day"],
    ["internal_marks", "Numeric", "Internal assessment marks (out of 100)"],
    ["participation", "Numeric (1-10)", "Class participation rating"],
    ["previous_score", "Numeric", "Previous academic performance (out of 100)"],
    ["math, science, english, computer score", "Numeric", "Subject-wise marks (used for subject analysis and recommendations)"],
    ["final_score", "Numeric", "Target for regression"],
    ["performance_level", "Derived", f"Target for classification: Low (< 50), Average (50 to 66.9), High (>= 67)"],
], widths=[5.2, 3.6, 7.2], font=9.5)
para("Subject scores and final score are not used as model inputs, which prevents target leakage.")

# ------------------------------------------------------------------------------------------ 6 prep
doc.add_heading("6. Data Preprocessing (Module 1)", 1)
mb = M["missing_before"]
table(["Step", "Method", "Result"], [
    ["Import dataset", "pandas.read_csv", f"{M['raw_rows']:,} rows"],
    ["Remove duplicates", "drop_duplicates()", f"{M['duplicates_before']} duplicate rows removed -> {M['clean_rows']:,} rows"],
    ["Missing values", "Median (numeric), mode (categorical)",
     "Filled: " + ", ".join(f"{k} ({v})" for k, v in mb.items()) + f". Missing after cleaning: {M['missing_after']}"],
    ["Encode categorical data", "One-hot encoding (pandas.get_dummies)", "Categorical columns converted to 0/1 columns"],
    ["Normalize features", "StandardScaler inside the sklearn Pipeline", "Scaling is fitted on training data only, avoiding leakage"],
    ["Create target classes", "Score cutoffs", f"Low {M['level_counts']['Low']}, Average {M['level_counts']['Average']}, High {M['level_counts']['High']}"],
], widths=[3.6, 5.4, 7], font=9.5)
para("Median imputation was chosen for numeric columns because it is not affected by outliers, and mode "
     "imputation for categorical columns because it keeps valid category values.")

# ------------------------------------------------------------------------------------------ 7 eda
doc.add_heading("7. Exploratory Data Analysis (Module 2)", 1)
ss = M["score_summary"]
para(f"The average final score is {ss['mean']:.1f} (standard deviation {ss['std']:.1f}), ranging from "
     f"{ss['min']:.1f} to {ss['max']:.1f}. Average performance is the largest group, with Low and High groups "
     "of comparable size, giving a reasonably balanced classification problem.")
figure("01_level_distribution.png", "Performance level distribution", 9)
figure("02_score_distribution.png", "Distribution of final scores with class cutoffs", 10)

doc.add_heading("7.1 Factors affecting performance", 2)
cs = M["corr_with_score"]
para("Correlation with the final score: " + ", ".join(f"{k.replace('_', ' ')} ({v:.2f})" for k, v in cs.items()) +
     ". All five factors are strongly positively related to the final score, with internal marks and previous "
     "academic performance the strongest.")
figure("03_correlation_heatmap.png", "Correlation heatmap of the numeric factors", 11)
figure("04_factors_vs_score.png", "Each factor plotted against the final score", 15.5)

gm = M["group_means"]
table(["Level", "Attendance %", "Study hrs/day", "Internal marks", "Participation", "Previous score"],
      [[lvl, f"{gm[lvl]['attendance']:.1f}", f"{gm[lvl]['study_hours']:.1f}", f"{gm[lvl]['internal_marks']:.1f}",
        f"{gm[lvl]['participation']:.1f}", f"{gm[lvl]['previous_score']:.1f}"] for lvl in ("Low", "Average", "High")],
      widths=[2.5, 2.7, 2.7, 2.8, 2.7, 2.8], font=9.5)
para("High performers attend about 20 percentage points more classes and study roughly 3 hours more per day than "
     "low performers.")
figure("05_boxplots_by_level.png", "Attendance, study hours and participation by performance level", 15.5)

doc.add_heading("7.2 Subject-wise analysis", 2)
sa = M["subject_avg_overall"]
para("Average subject scores: " + ", ".join(f"{k.replace('_score', '').title()} {v:.1f}" for k, v in sa.items()) +
     ". Mathematics has the lowest average and Computer the highest. Students in every level show the same "
     "ordering of subjects.")
figure("06_subject_wise.png", "Average subject score by performance level", 11)

doc.add_heading("7.3 Background factors", 2)
ce = M["cat_effects"]
para(f"Background factors have a much smaller effect than academic behaviour. Students with tutoring average "
     f"{ce['tutoring']['Yes']:.1f} versus {ce['tutoring']['No']:.1f} without, and students with internet access "
     f"average {ce['internet_access']['Yes']:.1f} versus {ce['internet_access']['No']:.1f}. Differences by gender "
     "are negligible. These small gaps are consistent with how the data was constructed and are not strong "
     "evidence about real-world effects.")
figure("07_categorical_factors.png", "Average final score by tutoring, internet access and parental education", 15.5)

# ------------------------------------------------------------------------------------------ 8 model
doc.add_heading("8. Model Development (Module 3)", 1)
para(f"The cleaned data was split 80% / 20% into {M['train_size']} training and {M['test_size']} test records "
     "(stratified by performance level, random_state = 42). Two prediction tasks were solved with the suggested "
     "algorithms:")
bullets(["Regression - predict the final score (Linear Regression, Decision Tree, Random Forest).",
         "Classification - predict Low / Average / High (Decision Tree, Random Forest). Linear Regression is a "
         "regression algorithm, so it is used for the score task only."])
para("Each model sits inside a scikit-learn Pipeline (StandardScaler then estimator). Tree depth was limited "
     "(Decision Tree max_depth 6; Random Forest 300 trees, max_depth 10) to reduce over-fitting. Models were "
     "compared on the held-out test set and with 5-fold cross-validation.")

# ------------------------------------------------------------------------------------------ 9 results
doc.add_heading("9. Results and Evaluation", 1)
doc.add_heading("9.1 Regression (final score)", 2)
table(["Model", "MAE", "RMSE", "R2 (test)", "CV R2 (mean +/- sd)"],
      [[k, v["MAE"], v["RMSE"], v["R2"], f"{v['CV_R2_mean']:.3f} +/- {v['CV_R2_std']:.3f}"]
       for k, v in M["regression"].items()], widths=[4, 2.4, 2.4, 2.6, 4.6])
br = M["best_regressor"]
para(f"{br} gave the best score prediction: the average error is about {M['regression'][br]['MAE']:.1f} marks "
     f"(R2 = {M['regression'][br]['R2']:.2f}). It beating the tree models is expected: the synthetic final score was "
     "generated mostly as a linear combination of the factors. On real data the non-linear Random Forest could "
     "well perform better, which is why all three models are kept in the code for comparison.")
figure("09_actual_vs_predicted.png", "Actual versus predicted final score on the test set", 9)

doc.add_heading("9.2 Classification (performance level)", 2)
table(["Model", "Accuracy (test)", "Macro F1", "CV accuracy (mean +/- sd)"],
      [[k, f"{v['Accuracy']:.3f}", f"{v['F1_macro']:.3f}", f"{v['CV_Accuracy_mean']:.3f} +/- {v['CV_Accuracy_std']:.3f}"]
       for k, v in M["classification"].items()], widths=[4.5, 3.5, 3.2, 5])
bc = M["best_classifier"]
para(f"{bc} is the better classifier ({M['classification'][bc]['Accuracy'] * 100:.1f}% test accuracy). Per-class "
     "results for this model:")
cr = M["classification_report"]
table(["Level", "Precision", "Recall", "F1-score", "Support"],
      [[l, f"{cr[l]['precision']:.2f}", f"{cr[l]['recall']:.2f}", f"{cr[l]['f1-score']:.2f}", int(cr[l]['support'])]
       for l in ("Low", "Average", "High")], widths=[3.5, 3, 3, 3, 3])
cm = M["confusion_matrix"]
para(f"The confusion matrix shows that nearly all mistakes are between neighbouring levels. No Low student was "
     f"predicted as High, and no High student as Low ({cm[0][2]} and {cm[2][0]} cases). This is the behaviour "
     "wanted from an early-warning tool.")
figure("08_confusion_matrix.png", "Confusion matrix of the Random Forest classifier (test set)", 8)
figure("10_model_comparison.png", "Model comparison", 13)

doc.add_heading("9.3 Key factors", 2)
fi = M["feature_importance"]
top5 = list(fi.items())[:5]
para("The Random Forest importance ranking confirms the EDA: " +
     ", ".join(f"{k.replace('_', ' ')} ({v * 100:.1f}%)" for k, v in top5) +
     ". Background attributes contribute under 1% each.")
figure("11_feature_importance.png", "Top 10 factors affecting performance", 11)

# ------------------------------------------------------------------------------------------ 10 dashboard
doc.add_heading("10. Performance Dashboard (Module 4)", 1)
para("The dashboard is built with Streamlit (app.py) and has seven tabs:")
table(["Tab", "What it shows"], [
    ["Overview", "Key metrics, performance-level distribution, factor importance, correlation heatmap"],
    ["Predict a student", "Form to enter student details; shows predicted score, level, confidence and recommendations"],
    ["Student comparison", "Side-by-side comparison of up to 5 students; actual vs predicted scores"],
    ["Subject-wise analysis", "Average and distribution of subject scores, filterable by level"],
    ["Performance trends", "Selected factor against the final score, binned trend line, categorical effects"],
    ["Recommendations", "Recommendations for any existing student"],
    ["Model details", "Metrics of all models and the confusion matrix"],
], widths=[4.2, 11.8])
para("[Insert dashboard screenshots here: Overview tab, Predict a student tab, Student comparison tab.]",
     italic=True, color=RGBColor(0xB0, 0x30, 0x30))

# ------------------------------------------------------------------------------------------ 11 rec
doc.add_heading("11. Recommendation System (Module 5)", 1)
para("The recommendation engine is rule based. For each student it compares attendance, study hours, "
     "participation and internal marks with the median of the High-performing group, and looks at subject "
     "scores relative to the student's own average. Recommendations produced include increasing study hours, "
     "improving attendance, focusing on weak subjects and attending practice sessions.")
doc.add_heading("Examples", 3)
for s in M["sample_recommendations"]:
    para(f"Student {s['student_id']} - predicted {s['predicted_level']} ({s['predicted_score']:.1f}); "
         f"actual score {s['actual_score']:.1f}", bold=True)
    bullets(s["recommendations"])

# ------------------------------------------------------------------------------------------ 12 testing
doc.add_heading("12. Testing", 1)
table(["Test", "Expected", "Outcome"], [
    ["Raw data contains duplicates and missing values", "Cleaned dataset has none", "Passed: 0 missing, 25 duplicates removed"],
    ["Prediction for one student entered manually", "Same encoded columns as training", "Passed: columns aligned with reindex"],
    ["Score output range", "Between 0 and 100", "Passed: clipped to 0-100"],
    ["High performer receives few recommendations", "Short list or praise", "Passed"],
    ["Low performer receives targeted advice", "Attendance, study hours, weak subjects", "Passed"],
    ["Dashboard loads without pre-built files", "Rebuilds data and model on first run", "Implemented in load_artifacts()"],
], widths=[5.8, 5, 5.2], font=9.5)
para("The pipeline and recommendation functions were run end to end during development. The Streamlit "
     "interface itself should be checked by running streamlit run app.py and exercising each tab before "
     "submission.")

# ------------------------------------------------------------------------------------------ 13 deploy
doc.add_heading("13. Deployment", 1)
numbered(["Push the project folder to a GitHub repository.",
          "Open Streamlit Community Cloud and choose New app.",
          "Select the repository, branch and app.py, then deploy.",
          "Copy the public URL into the report and presentation."])
para("Deployment link: [Paste link here]", bold=True)

# ------------------------------------------------------------------------------------------ 14 limits
doc.add_heading("14. Limitations and Future Work", 1)
bullets(["The dataset is synthetic. The pipeline should be re-run on a real Kaggle or institutional dataset before drawing conclusions about real students.",
         "The recommendation engine is rule based rather than learned.",
         "Predictions are decision support for teachers and students and must not be used as a final judgement of a student.",
         "Future work: add real datasets, SHAP-based per-student explanations, hyper-parameter tuning, and a database (MySQL) for storing students over time."])
para("As stated in the project scope, deep learning, real-time monitoring and large-scale platforms were "
     "intentionally left out.")

doc.add_heading("15. Conclusion", 1)
para(f"The project delivers a complete, working Student Performance Prediction System. Data cleaning and "
     f"preprocessing produced {M['clean_rows']:,} clean records; exploratory analysis showed that internal marks, "
     f"previous scores, study hours, attendance and participation are the main drivers of performance; the "
     f"{br} model predicts the final score with an average error of {M['regression'][br]['MAE']:.1f} marks, and the "
     f"{bc} classifies students into Low, Average and High with {M['classification'][bc]['Accuracy'] * 100:.1f}% accuracy. "
     "The Streamlit dashboard and the recommendation system turn these predictions into insights that a "
     "student or teacher can act on.")

doc.add_heading("16. References", 1)
bullets(["Pedregosa et al., Scikit-learn: Machine Learning in Python, JMLR 12, 2011.",
         "Streamlit documentation: https://docs.streamlit.io",
         "pandas documentation: https://pandas.pydata.org/docs",
         "Cortez, P. and Silva, A., Using Data Mining to Predict Secondary School Student Performance, 2008 (UCI Student Performance dataset).",
         "SkillOrbit Machine Learning Capstone Project brief."])

doc.add_heading("Appendix: Project Structure", 1)
table(["File", "Purpose"], [
    ["generate_dataset.py", "Creates the dataset"],
    ["ml_core.py", "Cleaning, features, prediction and recommendation logic"],
    ["train_model.py", "Preprocessing, EDA, training, evaluation, saving"],
    ["app.py", "Streamlit dashboard"],
    ["requirements.txt", "Python dependencies"],
    ["outputs/metrics.json", "All metrics quoted in this report"],
], widths=[5, 11])

# footer page numbers
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = fp.add_run("Student Performance Prediction System  |  Page ")
run.font.size = Pt(9)
r2 = fp.add_run()
for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        el = OxmlElement("w:fldChar")
        el.set(qn("w:fldCharType"), tag)
    else:
        el = OxmlElement("w:instrText")
        el.set(qn("xml:space"), "preserve")
        el.text = txt
    r2._r.append(el)
r2.font.size = Pt(9)

doc.save("Student_Performance_Prediction_Report.docx")
print("saved")
